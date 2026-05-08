#!/usr/bin/env python3
import json
import os
import re
import shutil
import subprocess
import sys
import time
import urllib.error
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Optional

APP_NAME = "GeminiSpeechToText"
TRANSCRIPTION_PROVIDERS = ("gemini", "whisper_local")
WHISPER_MODELS = ("tiny", "base", "small", "medium", "large", "turbo")
DEFAULT_ENV_VALUES = {
    "TRANSCRIPTION_PROVIDER": "gemini",
    "GEMINI_API_KEY": "",
    "GEMINI_MODEL": "gemini-2.5-pro",
    "MAX_RETRIES": "15",
    "RETRY_DELAY": "5",
    "BACKOFF_MAX": "300",
    "SEGMENT_DURATION_MINUTES": "15",
    "WHISPER_MODEL": "base",
    "WHISPER_LANGUAGE": "pt",
    "GEMINI_MODELS": (
        "gemini-3.1-pro-preview,gemini-3.1-flash-lite-preview,gemini-2.5-pro,"
        "gemini-2.0-flash,gemini-1.5-pro,gemini-1.5-flash,gemini-1.5-flash-8b"
    ),
}

VIDEO_EXTENSIONS = [".mp4", ".avi", ".mov", ".mkv", ".wmv", ".flv", ".webm", ".m4v", ".3gp", ".ogv"]
AUDIO_EXTENSIONS = [".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg"]
ALL_EXTENSIONS = VIDEO_EXTENSIONS + AUDIO_EXTENSIONS
EXTERNAL_BINARY_DIRNAME = "ffmpeg-bin"
_EXTERNAL_BINARY_CACHE: dict[str, str] = {}
_EXTERNAL_ENV_CONFIGURED = False


class ProcessingCancelled(Exception):
    pass


class HardQuotaExceeded(Exception):
    pass


@dataclass
class RuntimePaths:
    base_dir: Path
    source_dir: Path
    segments_dir: Path
    output_dir: Path
    env_file: Path


@dataclass
class PipelineSettings:
    transcription_provider: str
    api_key: str
    models: list[str]
    max_retries: int
    retry_delay: int
    backoff_max: int
    segment_duration_minutes: int
    whisper_model: str
    whisper_language: str


class PipelineContext:
    def __init__(
        self,
        logger: Optional[Callable[[str], None]] = None,
        stop_event=None,
        paths: Optional[RuntimePaths] = None,
        settings: Optional[PipelineSettings] = None,
        require_api_key: Optional[bool] = None,
    ):
        self.logger = logger or print
        self.stop_event = stop_event
        self.paths = ensure_runtime_dirs(paths or get_runtime_paths())
        self.settings = settings or load_settings(self.paths.env_file)
        self.client = None
        self.model_stats = {}
        self.whisper_model_instance = None

        if require_api_key is None:
            require_api_key = uses_gemini(self.settings)

        if require_api_key:
            if not has_configured_api_key(self.settings.api_key):
                raise ValueError(
                    f"GEMINI_API_KEY nao configurada no arquivo {self.paths.env_file}"
                )

            try:
                import google.genai as genai
            except ModuleNotFoundError as error:
                raise ModuleNotFoundError(
                    "Dependencia ausente: google-genai. Instale com 'python3 -m pip install google-genai'."
                ) from error

            self.client = genai.Client(api_key=self.settings.api_key)
            self.model_stats = {
                model: {"failures": 0, "last_503": 0}
                for model in self.settings.models
            }

    def log(self, message: str):
        self.logger(message)

    def check_cancelled(self):
        if self.stop_event is not None and self.stop_event.is_set():
            raise ProcessingCancelled("Processamento cancelado pelo usuario.")


def get_project_dir() -> Path:
    return Path(__file__).resolve().parent


def get_app_data_dir() -> Path:
    home = Path.home()

    if sys.platform.startswith("win"):
        root = Path(os.getenv("APPDATA", home / "AppData" / "Roaming"))
    elif sys.platform == "darwin":
        root = home / "Library" / "Application Support"
    else:
        root = Path(os.getenv("XDG_DATA_HOME", home / ".local" / "share"))

    return root / APP_NAME


def get_base_dir() -> Path:
    if getattr(sys, "frozen", False):
        return get_app_data_dir()
    return get_project_dir()


def get_runtime_paths(base_dir: Optional[Path] = None) -> RuntimePaths:
    runtime_dir = Path(base_dir or get_base_dir())
    return RuntimePaths(
        base_dir=runtime_dir,
        source_dir=runtime_dir / "arquivo",
        segments_dir=runtime_dir / "mp3",
        output_dir=runtime_dir / "trascricao",
        env_file=runtime_dir / ".env",
    )


def ensure_runtime_dirs(paths: RuntimePaths) -> RuntimePaths:
    paths.base_dir.mkdir(parents=True, exist_ok=True)
    paths.source_dir.mkdir(exist_ok=True)
    paths.segments_dir.mkdir(exist_ok=True)
    paths.output_dir.mkdir(exist_ok=True)
    ensure_env_file(paths.env_file)
    return paths


def render_env_template() -> str:
    return "\n".join(
        [
            "# Configuracao do provedor de transcricao",
            f"TRANSCRIPTION_PROVIDER={DEFAULT_ENV_VALUES['TRANSCRIPTION_PROVIDER']}",
            f"GEMINI_API_KEY={DEFAULT_ENV_VALUES['GEMINI_API_KEY']}",
            f"GEMINI_MODEL={DEFAULT_ENV_VALUES['GEMINI_MODEL']}",
            f"WHISPER_MODEL={DEFAULT_ENV_VALUES['WHISPER_MODEL']}",
            f"WHISPER_LANGUAGE={DEFAULT_ENV_VALUES['WHISPER_LANGUAGE']}",
            "# Configuracoes de resiliencia",
            f"MAX_RETRIES={DEFAULT_ENV_VALUES['MAX_RETRIES']}",
            f"RETRY_DELAY={DEFAULT_ENV_VALUES['RETRY_DELAY']}",
            f"BACKOFF_MAX={DEFAULT_ENV_VALUES['BACKOFF_MAX']}",
            "# Duracao dos segmentos (em minutos)",
            f"SEGMENT_DURATION_MINUTES={DEFAULT_ENV_VALUES['SEGMENT_DURATION_MINUTES']}",
            "# Modelos fallback (ordem de preferencia)",
            f"GEMINI_MODELS={DEFAULT_ENV_VALUES['GEMINI_MODELS']}",
            "",
        ]
    )


def ensure_env_file(env_path: Optional[Path] = None) -> Path:
    target = Path(env_path or get_runtime_paths().env_file)
    target.parent.mkdir(parents=True, exist_ok=True)

    if not target.exists():
        target.write_text(render_env_template(), encoding="utf-8")

    current_values = parse_env_file(target)
    for key, value in DEFAULT_ENV_VALUES.items():
        if current_values.get(key) is None:
            update_env_file_value(target, key, value)

    return target


def save_api_key(api_key: str, env_path: Optional[Path] = None) -> Path:
    target = ensure_env_file(env_path)
    update_env_file_value(target, "GEMINI_API_KEY", api_key.strip())
    return target


def save_settings(values: dict[str, str], env_path: Optional[Path] = None) -> Path:
    target = ensure_env_file(env_path)
    for key, value in values.items():
        update_env_file_value(target, key, str(value).strip())
    return target


def read_env_values(env_path: Optional[Path] = None) -> dict[str, str]:
    target = ensure_env_file(env_path)
    values = dict(DEFAULT_ENV_VALUES)
    loaded_values = parse_env_file(target)

    for key, default_value in DEFAULT_ENV_VALUES.items():
        env_value = os.getenv(key)
        if env_value is not None:
            values[key] = env_value
            continue

        loaded_value = loaded_values.get(key)
        if loaded_value is None:
            values[key] = default_value
        else:
            values[key] = loaded_value

    return values


def parse_env_file(env_path: Path) -> dict[str, str]:
    values = {}
    if not env_path.exists():
        return values

    for raw_line in env_path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue

        key, value = line.split("=", 1)
        values[key.strip()] = value.strip().strip('"').strip("'")

    return values


def update_env_file_value(env_path: Path, key: str, value: str):
    lines = []
    found = False

    if env_path.exists():
        lines = env_path.read_text(encoding="utf-8").splitlines()

    updated_lines = []
    for raw_line in lines:
        stripped = raw_line.strip()
        if stripped and not stripped.startswith("#") and "=" in raw_line:
            current_key = raw_line.split("=", 1)[0].strip()
            if current_key == key:
                updated_lines.append(f"{key}={value}")
                found = True
                continue

        updated_lines.append(raw_line)

    if not found:
        updated_lines.append(f"{key}={value}")

    env_path.write_text("\n".join(updated_lines).rstrip() + "\n", encoding="utf-8")


def load_settings(env_path: Optional[Path] = None) -> PipelineSettings:
    values = read_env_values(env_path)
    provider = values["TRANSCRIPTION_PROVIDER"].strip().lower() or DEFAULT_ENV_VALUES["TRANSCRIPTION_PROVIDER"]
    if provider not in TRANSCRIPTION_PROVIDERS:
        provider = DEFAULT_ENV_VALUES["TRANSCRIPTION_PROVIDER"]

    models = [model.strip() for model in values["GEMINI_MODELS"].split(",") if model.strip()]
    whisper_model = values["WHISPER_MODEL"].strip() or DEFAULT_ENV_VALUES["WHISPER_MODEL"]
    if whisper_model not in WHISPER_MODELS:
        whisper_model = DEFAULT_ENV_VALUES["WHISPER_MODEL"]
    whisper_language = normalize_whisper_language(values["WHISPER_LANGUAGE"])

    return PipelineSettings(
        transcription_provider=provider,
        api_key=values["GEMINI_API_KEY"].strip(),
        models=models,
        max_retries=int(values["MAX_RETRIES"]),
        retry_delay=int(values["RETRY_DELAY"]),
        backoff_max=int(values["BACKOFF_MAX"]),
        segment_duration_minutes=int(values["SEGMENT_DURATION_MINUTES"]),
        whisper_model=whisper_model,
        whisper_language=whisper_language,
    )


def normalize_whisper_language(language: str) -> str:
    cleaned = (language or "").strip().lower().replace("_", "-")
    if not cleaned:
        return ""

    try:
        from whisper.tokenizer import LANGUAGES, TO_LANGUAGE_CODE
    except Exception:
        return cleaned.split("-", 1)[0]

    if cleaned in LANGUAGES:
        return cleaned
    if cleaned in TO_LANGUAGE_CODE:
        return TO_LANGUAGE_CODE[cleaned]

    base_language = cleaned.split("-", 1)[0]
    if base_language in LANGUAGES:
        return base_language
    if base_language in TO_LANGUAGE_CODE:
        return TO_LANGUAGE_CODE[base_language]

    return cleaned


def uses_gemini(settings: PipelineSettings) -> bool:
    return settings.transcription_provider == "gemini"


def uses_whisper(settings: PipelineSettings) -> bool:
    return settings.transcription_provider == "whisper_local"


def has_configured_api_key(api_key: str) -> bool:
    cleaned_key = (api_key or "").strip()
    return bool(cleaned_key and cleaned_key != "sua_chave_api_gemini_aqui")


def sanitize_filename(filename: str) -> str:
    path = Path(filename)
    suffix = path.suffix
    stem = path.stem

    normalized = unicodedata.normalize("NFKD", stem)
    ascii_stem = normalized.encode("ascii", "ignore").decode("ascii")
    ascii_stem = re.sub(r"[^A-Za-z0-9._-]+", "_", ascii_stem)
    ascii_stem = re.sub(r"_+", "_", ascii_stem).strip("._-")

    if not ascii_stem:
        ascii_stem = "arquivo"

    return f"{ascii_stem}{suffix}"


def generate_unique_path(directory: Path, filename: str) -> Path:
    candidate = directory / filename
    if not candidate.exists():
        return candidate

    path = Path(filename)
    stem = path.stem
    suffix = path.suffix
    counter = 1

    while True:
        candidate = directory / f"{stem}_{counter}{suffix}"
        if not candidate.exists():
            return candidate
        counter += 1


def is_media_file(filename: str) -> bool:
    return Path(filename).suffix.lower() in ALL_EXTENSIONS


def is_video_file(filename: str) -> bool:
    return Path(filename).suffix.lower() in VIDEO_EXTENSIONS


def sleep_with_cancel(context: PipelineContext, seconds: int):
    if seconds <= 0:
        return

    deadline = time.time() + seconds
    while time.time() < deadline:
        context.check_cancelled()
        time.sleep(min(0.5, deadline - time.time()))


def prepend_env_path(env_name: str, directory: Path):
    if not directory.exists():
        return

    current_entries = [entry for entry in os.getenv(env_name, "").split(os.pathsep) if entry]
    directory_str = str(directory)
    if directory_str in current_entries:
        return

    os.environ[env_name] = os.pathsep.join([directory_str, *current_entries]) if current_entries else directory_str


def iter_external_binary_dirs() -> list[Path]:
    candidates: list[Path] = []

    if getattr(sys, "frozen", False):
        executable_dir = Path(sys.executable).resolve().parent
        candidates.extend(
            [
                executable_dir / EXTERNAL_BINARY_DIRNAME,
                executable_dir,
            ]
        )

        if sys.platform == "darwin":
            contents_dir = executable_dir.parent
            candidates.extend(
                [
                    contents_dir / "Frameworks" / EXTERNAL_BINARY_DIRNAME,
                    contents_dir / "Resources" / EXTERNAL_BINARY_DIRNAME,
                    contents_dir / "MacOS" / EXTERNAL_BINARY_DIRNAME,
                ]
            )

        bundle_root = getattr(sys, "_MEIPASS", None)
        if bundle_root:
            bundle_dir = Path(bundle_root)
            candidates.extend(
                [
                    bundle_dir / EXTERNAL_BINARY_DIRNAME,
                    bundle_dir,
                ]
            )
    else:
        project_dir = get_project_dir()
        candidates.extend(
            [
                project_dir / EXTERNAL_BINARY_DIRNAME,
                project_dir / "bin",
            ]
        )

    unique_candidates = []
    seen = set()
    for candidate in candidates:
        resolved = candidate.resolve()
        if resolved in seen:
            continue
        seen.add(resolved)
        unique_candidates.append(resolved)

    return unique_candidates


def configure_external_runtime_environment():
    global _EXTERNAL_ENV_CONFIGURED

    if _EXTERNAL_ENV_CONFIGURED:
        return

    for directory in iter_external_binary_dirs():
        if not directory.exists():
            continue

        prepend_env_path("PATH", directory)

        lib_dir = directory / "lib"
        if not lib_dir.exists():
            continue

        if sys.platform == "darwin":
            prepend_env_path("DYLD_LIBRARY_PATH", lib_dir)
        elif not sys.platform.startswith("win"):
            prepend_env_path("LD_LIBRARY_PATH", lib_dir)

    _EXTERNAL_ENV_CONFIGURED = True


def get_external_binary_filename(binary_name: str) -> str:
    if sys.platform.startswith("win") and not binary_name.lower().endswith(".exe"):
        return f"{binary_name}.exe"
    return binary_name


def resolve_external_binary(binary_name: str) -> Optional[str]:
    cached = _EXTERNAL_BINARY_CACHE.get(binary_name)
    if cached and Path(cached).exists():
        return cached

    configure_external_runtime_environment()

    filename = get_external_binary_filename(binary_name)
    for directory in iter_external_binary_dirs():
        candidate = directory / filename
        if candidate.exists():
            _EXTERNAL_BINARY_CACHE[binary_name] = str(candidate)
            prepend_env_path("PATH", candidate.parent)
            return str(candidate)

    resolved = shutil.which(binary_name) or shutil.which(filename)
    if resolved:
        _EXTERNAL_BINARY_CACHE[binary_name] = resolved
        prepend_env_path("PATH", Path(resolved).parent)
        return resolved

    return None


def ensure_external_dependencies():
    missing = [binary for binary in ("ffmpeg", "ffprobe") if not resolve_external_binary(binary)]
    if missing:
        raise FileNotFoundError(
            "Dependencias ausentes: " + ", ".join(missing) + ". Instale o FFmpeg e garanta que ele esteja no PATH."
        )


def ensure_whisper_dependency():
    try:
        import whisper  # noqa: F401
    except ModuleNotFoundError as error:
        raise ModuleNotFoundError(
            "Dependencia ausente: openai-whisper. Instale com 'python3 -m pip install openai-whisper'."
        ) from error
    except Exception as error:
        message = str(error)
        if "Numpy is not available" in message or "_ARRAY_API not found" in message:
            raise RuntimeError(
                "Ambiente Whisper incompatível: o PyTorch instalado nao funciona com a versao atual do NumPy. "
                "No mesmo ambiente da interface, execute: "
                "\"python3 -m pip install 'numpy<2' --force-reinstall\"."
            ) from error
        raise


def shutil_which(binary_name: str) -> Optional[str]:
    return resolve_external_binary(binary_name)


def convert_to_mp3(video_file: str, output_file: str):
    cmd = [
        resolve_external_binary("ffmpeg") or "ffmpeg",
        "-i",
        video_file,
        "-vn",
        "-c:a",
        "libmp3lame",
        "-b:a",
        "128k",
        "-ar",
        "44100",
        "-f",
        "mp3",
        "-avoid_negative_ts",
        "make_zero",
        "-y",
        output_file,
    ]
    result = subprocess.run(cmd, capture_output=True, text=True, check=True)

    if not os.path.exists(output_file) or os.path.getsize(output_file) == 0:
        raise RuntimeError("Conversao falhou: arquivo nao criado ou vazio.")

    return result


def get_audio_duration(audio_file: str) -> float:
    cmd = [
        resolve_external_binary("ffprobe") or "ffprobe",
        "-v",
        "quiet",
        "-show_entries",
        "format=duration",
        "-of",
        "csv=p=0",
        audio_file,
    ]
    result = subprocess.run(cmd, capture_output=True, text=True, check=True)
    return float(result.stdout.strip())


def split_audio(audio_file: str, output_dir: str, segment_duration_minutes: int) -> list[str]:
    duration = get_audio_duration(audio_file)
    segment_duration = segment_duration_minutes * 60

    base_name = Path(audio_file).stem
    segments = []

    start_time = 0
    segment_num = 1

    while start_time < duration:
        remaining_time = duration - start_time
        current_segment_duration = min(segment_duration, remaining_time)
        output_file = os.path.join(output_dir, f"{base_name}-{segment_num:02d}.mp3")

        cmd = [
            resolve_external_binary("ffmpeg") or "ffmpeg",
            "-i",
            audio_file,
            "-ss",
            str(start_time),
            "-t",
            str(current_segment_duration),
            "-acodec",
            "copy",
            "-y",
            output_file,
        ]

        subprocess.run(cmd, check=True)
        segments.append(output_file)

        start_time += segment_duration
        segment_num += 1

    return segments


def process_media_files(context: PipelineContext) -> list[str]:
    context.log("ETAPA 1: Processando arquivos de midia...")

    normalize_source_filenames(context)

    media_files = sorted(
        [
            path
            for path in context.paths.source_dir.iterdir()
            if path.is_file() and is_media_file(path.name)
        ]
    )

    if not media_files:
        raise FileNotFoundError(
            f"Nenhum arquivo de midia encontrado em {context.paths.source_dir}"
        )

    generated_segments = []

    for media_file in media_files:
        context.check_cancelled()
        context.log(f"Processando: {media_file.name}")
        base_name = media_file.stem

        try:
            if is_video_file(media_file.name):
                context.log("  Convertendo video para MP3...")
                temp_mp3 = context.paths.segments_dir / f"{base_name}_temp.mp3"
                convert_to_mp3(str(media_file), str(temp_mp3))
                audio_to_split = temp_mp3
            else:
                context.log("  Arquivo ja e audio, pulando conversao...")
                audio_to_split = media_file

            context.log(
                f"  Dividindo em segmentos de {context.settings.segment_duration_minutes} minutos..."
            )
            segments = split_audio(
                str(audio_to_split),
                str(context.paths.segments_dir),
                context.settings.segment_duration_minutes,
            )
            generated_segments.extend(segments)

            if is_video_file(media_file.name) and audio_to_split.exists():
                audio_to_split.unlink()

            context.log(f"  OK - Criados {len(segments)} segmentos")
        except subprocess.CalledProcessError as error:
            raise RuntimeError(f"Erro ao processar {media_file.name}: {error}") from error
        except Exception as error:
            raise RuntimeError(f"Erro inesperado ao processar {media_file.name}: {error}") from error

    return generated_segments


def normalize_source_filenames(context: PipelineContext):
    for source_file in context.paths.source_dir.iterdir():
        if not source_file.is_file() or not is_media_file(source_file.name):
            continue

        sanitized_name = sanitize_filename(source_file.name)
        if sanitized_name == source_file.name:
            continue

        destination = generate_unique_path(context.paths.source_dir, sanitized_name)
        source_file.rename(destination)
        context.log(
            f"Arquivo renomeado para evitar erro de encoding: {source_file.name} -> {destination.name}"
        )


def is_retryable_error(error: Exception) -> tuple[bool, int]:
    error_str = str(error).lower()
    if "503" in error_str or "unavailable" in error_str or "high demand" in error_str:
        return True, 503
    if is_hard_quota_error(error_str):
        return False, 429
    if "429" in error_str or "quota" in error_str or "rate limit" in error_str:
        return True, 429
    return False, 0


def is_hard_quota_error(error_text: str) -> bool:
    return "limit: 0" in error_text or "perday" in error_text


def build_quota_help_message(model: str) -> str:
    return (
        f"Cota indisponivel para o modelo {model}. "
        "A chave atual nao tem limite utilizavel nesse modelo ou a cota diaria do projeto esta zerada. "
        "Troque para outra API key/projeto com billing e quota ativa, ou ajuste GEMINI_MODELS para usar apenas modelos com cota disponivel."
    )


def format_timestamp(seconds: float) -> str:
    total_seconds = max(0, int(seconds))
    hours, remainder = divmod(total_seconds, 3600)
    minutes, secs = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def get_whisper_model(context: PipelineContext):
    if context.whisper_model_instance is not None:
        return context.whisper_model_instance

    ensure_whisper_dependency()

    try:
        import whisper
    except ModuleNotFoundError as error:
        raise ModuleNotFoundError(
            "Dependencia ausente: openai-whisper. Instale com 'python3 -m pip install openai-whisper'."
        ) from error

    context.log(f"Carregando modelo Whisper local: {context.settings.whisper_model}")
    try:
        context.whisper_model_instance = whisper.load_model(context.settings.whisper_model)
    except urllib.error.URLError as error:
        raise RuntimeError(
            "Nao foi possivel baixar o modelo Whisper. Verifique sua conexao com a internet "
            "e tente novamente."
        ) from error
    return context.whisper_model_instance


def upload_audio_file(file_path: str, context: PipelineContext) -> str:
    context.log(f"    Fazendo upload de {os.path.basename(file_path)}...")

    for attempt in range(context.settings.max_retries):
        context.check_cancelled()
        try:
            upload_response = context.client.files.upload(
                file=file_path,
                config={"display_name": sanitize_filename(os.path.basename(file_path))},
            )

            file_uri = upload_response.uri
            sleep_with_cancel(context, 5)
            return file_uri
        except Exception as error:
            is_retryable, _ = is_retryable_error(error)
            context.log(f"    AVISO - Upload tentativa {attempt + 1} falhou: {error}")

            if not is_retryable or attempt == context.settings.max_retries - 1:
                raise

            wait_time = min(
                context.settings.retry_delay * (2 ** attempt),
                context.settings.backoff_max,
            )
            context.log(f"    Aguardando {wait_time}s...")
            sleep_with_cancel(context, wait_time)

    raise RuntimeError("Falha inesperada no upload do arquivo.")


def identify_personas(first_audio_file: str, context: PipelineContext) -> dict:
    context.log("  Identificando personas no primeiro arquivo...")
    file_uri = upload_audio_file(first_audio_file, context)

    prompt = """
Analise este audio e identifique todas as pessoas que falam durante a conversa.

Para cada pessoa, forneca:
1. Um nome ou identificador simples
2. Caracteristicas distintivas da voz
3. Contexto ou papel na conversa, se aparente

Responda apenas em JSON valido:
{
  "personas": [
    {
      "id": "Pessoa1",
      "nome": "Nome ou identificador",
      "caracteristicas": "Descricao da voz e caracteristicas",
      "papel": "Papel na conversa se identificavel"
    }
  ]
}
"""

    for model in context.settings.models:
        for attempt in range(3):
            context.check_cancelled()
            try:
                context.log(f"    Modelo: {model} (tentativa {attempt + 1})")

                response = context.client.models.generate_content(
                    model=model,
                    contents=[
                        {
                            "parts": [
                                {"text": prompt},
                                {"file_data": {"file_uri": file_uri}},
                            ]
                        }
                    ],
                )

                response_text = response.candidates[0].content.parts[0].text
                if "```json" in response_text:
                    json_text = response_text.split("```json", 1)[1].split("```", 1)[0].strip()
                else:
                    json_text = response_text.strip()

                personas = json.loads(json_text)
                context.log(f"    OK - Personas identificadas com {model}")
                return personas
            except Exception as error:
                is_retryable, error_code = is_retryable_error(error)
                context.log(f"    Falha em {model}: {error}")

                if error_code == 429 and is_hard_quota_error(str(error).lower()):
                    raise HardQuotaExceeded(build_quota_help_message(model)) from error

                if error_code == 503:
                    context.model_stats[model]["last_503"] = time.time()
                    context.model_stats[model]["failures"] += 1
                    break

                if not is_retryable:
                    break

                if attempt < 2:
                    sleep_with_cancel(context, context.settings.retry_delay * (attempt + 1))

    context.log("    Todos os modelos falharam na identificacao de personas. Usando fallback.")
    return {
        "personas": [
            {
                "id": "Pessoa1",
                "nome": "Primeira pessoa",
                "caracteristicas": "Voz principal",
                "papel": "Participante",
            },
            {
                "id": "Pessoa2",
                "nome": "Segunda pessoa",
                "caracteristicas": "Voz secundaria",
                "papel": "Participante",
            },
        ]
    }


def transcribe_audio_with_personas(audio_file: str, personas_info: dict, context: PipelineContext) -> str:
    context.log(f"    Transcrevendo {os.path.basename(audio_file)}...")
    file_uri = upload_audio_file(audio_file, context)

    personas_list = "\n".join(
        [f"- {persona['id']}: {persona['caracteristicas']}" for persona in personas_info["personas"]]
    )

    prompts = [
        f"""
Transcreva este audio de forma completamente fiel, identificando cada fala por persona.

PERSONAS IDENTIFICADAS:
{personas_list}

INSTRUCOES:
1. Transcreva exatamente o que e falado
2. Identifique cada fala com a persona correspondente
3. Mantenha a ordem cronologica exata
4. Nao adicione comentarios ou resumos
5. Inclua pausas e hesitacoes quando existirem
6. Se nao conseguir identificar quem fala, use "Nao identificado"

FORMATO:
[PERSONA]: Fala
""",
        f"Transcreva este audio identificando falantes como [Pessoa1], [Pessoa2].\n\nPERSONAS:\n{personas_list}",
        "Faca a transcricao completa deste audio identificando diferentes falantes.",
        "Transcreva este audio completamente.",
    ]

    for strategy_idx, prompt in enumerate(prompts, 1):
        context.log(f"    Estrategia {strategy_idx}/{len(prompts)}")

        for model in context.settings.models:
            if time.time() - context.model_stats[model]["last_503"] < 600:
                continue

            for attempt in range(3):
                context.check_cancelled()
                try:
                    context.log(f"      Modelo: {model} (tentativa {attempt + 1})")

                    response = context.client.models.generate_content(
                        model=model,
                        contents=[
                            {
                                "parts": [
                                    {"text": prompt},
                                    {"file_data": {"file_uri": file_uri}},
                                ]
                            }
                        ],
                    )

                    transcription = response.candidates[0].content.parts[0].text
                    if len(transcription.strip()) < 10:
                        raise RuntimeError("Transcricao muito curta.")

                    context.log(f"      OK - Sucesso com {model}")
                    return transcription
                except Exception as error:
                    is_retryable, error_code = is_retryable_error(error)
                    context.log(f"      Falha em {model}: {error}")

                    if error_code == 429 and is_hard_quota_error(str(error).lower()):
                        raise HardQuotaExceeded(build_quota_help_message(model)) from error

                    if error_code == 503:
                        context.model_stats[model]["last_503"] = time.time()
                        context.model_stats[model]["failures"] += 1
                        break

                    if not is_retryable:
                        break

                    if attempt < 2:
                        sleep_with_cancel(context, context.settings.retry_delay * (attempt + 1))

        if strategy_idx < len(prompts):
            context.log("    Tentando proxima estrategia em 10s...")
            sleep_with_cancel(context, 10)

    return (
        f"[ERRO_CRITICO]: Falha total na transcricao de {os.path.basename(audio_file)} "
        f"apos {len(prompts)} estrategias e {len(context.settings.models)} modelos"
    )


def transcribe_audio_with_whisper(audio_file: str, context: PipelineContext) -> str:
    context.log(f"    Transcrevendo localmente com Whisper: {os.path.basename(audio_file)}...")
    model = get_whisper_model(context)

    options = {
        "task": "transcribe",
        "verbose": False,
        "fp16": False,
    }
    if context.settings.whisper_language:
        options["language"] = context.settings.whisper_language

    try:
        result = model.transcribe(audio_file, **options)
    except ValueError as error:
        message = str(error)
        if "Unsupported language:" in message:
            raise RuntimeError(
                "Idioma Whisper invalido. Use autodeteccao deixando o campo vazio, "
                "ou informe um codigo como `pt`, `en`, `es`. Variantes como `pt-BR` "
                "sao convertidas automaticamente para `pt`."
            ) from error
        raise
    detected_language = result.get("language")
    if detected_language:
        context.log(f"      Idioma detectado: {detected_language}")

    segments = result.get("segments") or []
    if not segments:
        text = (result.get("text") or "").strip()
        return text or "[SEM_TEXTO]: O Whisper nao retornou conteudo para este trecho."

    lines = []
    for segment in segments:
        segment_text = (segment.get("text") or "").strip()
        if not segment_text:
            continue

        start = format_timestamp(segment.get("start", 0))
        end = format_timestamp(segment.get("end", 0))
        lines.append(f"[{start} - {end}] {segment_text}")

    return "\n".join(lines)


def create_transcription_document(
    transcriptions: list[tuple[str, str]],
    personas_info: dict,
    output_path: Path,
    context: PipelineContext,
):
    context.log("  Criando documento Word com transcricoes...")
    try:
        from docx import Document
    except ModuleNotFoundError as error:
        raise ModuleNotFoundError(
            "Dependencia ausente: python-docx. Instale com 'python3 -m pip install python-docx'."
        ) from error

    document = Document()

    document.add_heading("Transcricao Completa do Audio", 0)
    document.add_heading("Personas Identificadas", level=1)

    for persona in personas_info["personas"]:
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{persona['id']}: ").bold = True
        paragraph.add_run(persona["nome"])
        if persona.get("caracteristicas"):
            paragraph.add_run(f" - {persona['caracteristicas']}")
        if persona.get("papel"):
            paragraph.add_run(f" ({persona['papel']})")

    document.add_page_break()

    for index, (filename, transcription) in enumerate(transcriptions, 1):
        document.add_heading(f"Segmento {index:02d} - {filename}", level=1)

        for line in transcription.strip().split("\n"):
            if not line.strip():
                continue

            if line.startswith("[") and "]:" in line:
                parts = line.split("]:", 1)
                if len(parts) == 2:
                    persona = parts[0] + "]:"
                    speech = parts[1].strip()

                    paragraph = document.add_paragraph()
                    paragraph.add_run(persona).bold = True
                    paragraph.add_run(f" {speech}")
                    continue

            document.add_paragraph(line)

        if index < len(transcriptions):
            document.add_paragraph("-" * 50)

    document.save(output_path)
    context.log(f"  OK - Documento salvo em: {output_path}")


def create_whisper_transcription_document(
    transcriptions: list[tuple[str, str]],
    output_path: Path,
    context: PipelineContext,
):
    context.log("  Criando documento Word com transcricoes locais do Whisper...")
    try:
        from docx import Document
    except ModuleNotFoundError as error:
        raise ModuleNotFoundError(
            "Dependencia ausente: python-docx. Instale com 'python3 -m pip install python-docx'."
        ) from error

    document = Document()
    document.add_heading("Transcricao Completa do Audio", 0)

    summary = document.add_paragraph()
    summary.add_run("Motor de transcricao: ").bold = True
    summary.add_run("Whisper local")

    model_line = document.add_paragraph()
    model_line.add_run("Modelo Whisper: ").bold = True
    model_line.add_run(context.settings.whisper_model)

    language_line = document.add_paragraph()
    language_line.add_run("Idioma configurado: ").bold = True
    language_line.add_run(context.settings.whisper_language or "auto")

    document.add_page_break()

    for index, (filename, transcription) in enumerate(transcriptions, 1):
        document.add_heading(f"Segmento {index:02d} - {filename}", level=1)

        for line in transcription.strip().split("\n"):
            if line.strip():
                document.add_paragraph(line)

        if index < len(transcriptions):
            document.add_paragraph("-" * 50)

    document.save(output_path)
    context.log(f"  OK - Documento salvo em: {output_path}")


def transcribe_audio_segments(context: PipelineContext) -> Path:
    context.log("ETAPA 2: Transcrevendo arquivos de audio...")

    mp3_files = sorted(context.paths.segments_dir.glob("*.mp3"))
    if not mp3_files:
        raise FileNotFoundError(
            f"Nenhum arquivo MP3 encontrado em {context.paths.segments_dir}"
        )

    context.log(f"  Encontrados {len(mp3_files)} arquivos de audio para transcrever")
    base_name = mp3_files[0].stem.replace("-01", "")

    if uses_whisper(context.settings):
        transcriptions = []
        for index, mp3_file in enumerate(mp3_files, 1):
            context.check_cancelled()
            context.log(f"  Processando arquivo {index}/{len(mp3_files)}")
            transcription = transcribe_audio_with_whisper(str(mp3_file), context)
            transcriptions.append((mp3_file.name, transcription))

        output_path = context.paths.output_dir / f"{base_name}_transcricao_whisper.docx"
        create_whisper_transcription_document(transcriptions, output_path, context)
    else:
        first_file = str(mp3_files[0])
        personas_info = identify_personas(first_file, context)

        context.log("  Personas identificadas:")
        for persona in personas_info["personas"]:
            context.log(f"    - {persona['id']}: {persona['nome']}")

        transcriptions = []
        for index, mp3_file in enumerate(mp3_files, 1):
            context.check_cancelled()
            context.log(f"  Processando arquivo {index}/{len(mp3_files)}")
            transcription = transcribe_audio_with_personas(str(mp3_file), personas_info, context)
            transcriptions.append((mp3_file.name, transcription))
            sleep_with_cancel(context, 3)

        output_path = context.paths.output_dir / f"{base_name}_transcricao_completa.docx"
        create_transcription_document(transcriptions, personas_info, output_path, context)

    context.log("Removendo arquivos MP3 temporarios...")
    for mp3_file in mp3_files:
        try:
            mp3_file.unlink()
            context.log(f"  OK - Removido: {mp3_file.name}")
        except Exception as error:
            context.log(f"  AVISO - Erro ao remover {mp3_file.name}: {error}")

    if context.paths.segments_dir.exists() and not any(context.paths.segments_dir.iterdir()):
        context.paths.segments_dir.rmdir()
        context.paths.segments_dir.mkdir(exist_ok=True)

    context.log("Transcricao completa finalizada com sucesso.")
    context.log(f"Documento final: {output_path}")
    return output_path


def run_pipeline(
    logger: Optional[Callable[[str], None]] = None,
    stop_event=None,
    paths: Optional[RuntimePaths] = None,
) -> Path:
    ensure_external_dependencies()
    context = PipelineContext(logger=logger, stop_event=stop_event, paths=paths)
    if uses_whisper(context.settings):
        ensure_whisper_dependency()

    context.log("INICIANDO PROCESSAMENTO COMPLETO")
    context.log("=" * 50)
    context.log(f"Provedor selecionado: {context.settings.transcription_provider}")

    process_media_files(context)
    context.log("Processamento de midia concluido.")
    output_path = transcribe_audio_segments(context)

    context.log("=" * 50)
    context.log("PROCESSAMENTO COMPLETO FINALIZADO COM SUCESSO")
    return output_path


def main():
    try:
        run_pipeline()
    except ProcessingCancelled as error:
        print(f"Processamento interrompido: {error}")
    except KeyboardInterrupt:
        print("Processamento interrompido pelo usuario.")
    except Exception as error:
        print(f"Erro fatal: {error}")
        raise


if __name__ == "__main__":
    main()
