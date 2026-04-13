#!/usr/bin/env python3
import os
import json
import time
import google.genai as genai
from pathlib import Path
from docx import Document
from docx.shared import Inches
import glob
from dotenv import load_dotenv

# Carregar variáveis de ambiente
load_dotenv()

# Configurar a API do Gemini
GEMINI_API_KEY = os.getenv('GEMINI_API_KEY')
GEMINI_MODEL = os.getenv('GEMINI_MODEL', 'gemini-2.5-pro')
if not GEMINI_API_KEY:
    raise ValueError("GEMINI_API_KEY não encontrada no arquivo .env")
client = genai.Client(api_key=GEMINI_API_KEY)

def upload_audio_file(file_path):
    """Upload do arquivo de áudio para o Gemini"""
    print(f"Fazendo upload de {os.path.basename(file_path)}...")
    
    # Faz upload do arquivo
    upload_response = client.files.upload(
        path=file_path,
        config={'display_name': os.path.basename(file_path)}
    )
    
    file_uri = upload_response.uri
    print(f"Arquivo carregado: {file_uri}")
    
    # Aguardar o processamento se necessário
    time.sleep(5)  # Pequena pausa para garantir processamento
    
    return file_uri

def identify_personas(first_audio_file):
    """Identifica as personas no primeiro arquivo de áudio"""
    print("Identificando personas no primeiro arquivo...")
    
    file_uri = upload_audio_file(first_audio_file)
    
    prompt = """
Analise este áudio e identifique todas as pessoas que falam durante a conversa.

Para cada pessoa, forneça:
1. Um nome/identificador simples (ex: Pessoa 1, Pessoa 2, ou se mencionado o nome real)
2. Características distintivas da voz (tom, velocidade, etc.)
3. Contexto/papel na conversa (se aparente)

Responda APENAS em formato JSON válido:
{
  "personas": [
    {
      "id": "Pessoa1",
      "nome": "Nome ou identificador",
      "caracteristicas": "Descrição da voz e características",
      "papel": "Papel na conversa se identificável"
    }
  ]
}

Seja preciso e objetivo. Não adicione texto explicativo fora do JSON.
"""
    
    try:
        response = client.models.generate_content(
            model=GEMINI_MODEL,
            contents=[
                {'parts': [
                    {'text': prompt},
                    {'file_data': {'file_uri': file_uri}}
                ]}
            ]
        )
        
        response_text = response.candidates[0].content.parts[0].text
        print(f"Resposta do Gemini: {response_text}")
        
        # Tentar extrair JSON da resposta
        if '```json' in response_text:
            json_text = response_text.split('```json')[1].split('```')[0].strip()
        else:
            json_text = response_text.strip()
            
        return json.loads(json_text)
        
    except Exception as e:
        print(f"Erro ao analisar personas: {e}")
        # Fallback: criar personas genéricas
        return {
            "personas": [
                {
                    "id": "Pessoa1",
                    "nome": "Primeira pessoa",
                    "caracteristicas": "Voz principal",
                    "papel": "Participante"
                },
                {
                    "id": "Pessoa2", 
                    "nome": "Segunda pessoa",
                    "caracteristicas": "Voz secundária",
                    "papel": "Participante"
                }
            ]
        }

def transcribe_audio_with_personas(audio_file, personas_info):
    """Transcreve um arquivo de áudio identificando cada persona"""
    print(f"Transcrevendo {os.path.basename(audio_file)}...")
    
    file_uri = upload_audio_file(audio_file)
    
    personas_list = "\n".join([
        f"- {p['id']}: {p['caracteristicas']}"
        for p in personas_info['personas']
    ])
    
    prompt = f"""
Transcreva este áudio de forma COMPLETAMENTE FIEL, identificando cada fala por persona.

PERSONAS IDENTIFICADAS:
{personas_list}

INSTRUÇÕES CRÍTICAS:
1. Transcreva EXATAMENTE o que é falado, sem omitir, resumir ou parafrasear
2. Identifique cada fala com a persona correspondente
3. Mantenha a ordem cronológica exata das falas
4. NÃO adicione comentários, interpretações ou textos complementares
5. Inclua pausas significativas, hesitações (eh, ah, hm) se presentes
6. Se não conseguir identificar claramente quem está falando, use "Não identificado"

FORMATO DA RESPOSTA:
[PERSONA]: Transcrição exata da fala

[PERSONA]: Transcrição exata da próxima fala

Seja absolutamente fiel ao áudio. Transcreva tudo que for dito.
"""
    
    try:
        response = client.models.generate_content(
            model=GEMINI_MODEL,
            contents=[
                {'parts': [
                    {'text': prompt},
                    {'file_data': {'file_uri': file_uri}}
                ]}
            ]
        )
        
        return response.candidates[0].content.parts[0].text
        
    except Exception as e:
        print(f"Erro ao transcrever {audio_file}: {e}")
        return f"[ERRO]: Falha na transcrição de {os.path.basename(audio_file)}"

def create_transcription_document(transcriptions, personas_info, output_path):
    """Cria um documento Word com todas as transcrições"""
    print("Criando documento Word com transcrições...")
    
    doc = Document()
    
    # Título
    title = doc.add_heading('Transcrição Completa do Áudio', 0)
    
    # Seção de Personas
    doc.add_heading('Personas Identificadas', level=1)
    for persona in personas_info['personas']:
        p = doc.add_paragraph()
        p.add_run(f"{persona['id']}: ").bold = True
        p.add_run(f"{persona['nome']}")
        if persona.get('caracteristicas'):
            p.add_run(f" - {persona['caracteristicas']}")
        if persona.get('papel'):
            p.add_run(f" ({persona['papel']})")
    
    doc.add_page_break()
    
    # Transcrições por segmento
    for i, (filename, transcription) in enumerate(transcriptions, 1):
        doc.add_heading(f'Segmento {i:02d} - {filename}', level=1)
        
        # Dividir por falas e formatar
        lines = transcription.strip().split('\n')
        for line in lines:
            if line.strip():
                if line.startswith('[') and ']:' in line:
                    # É uma fala identificada
                    parts = line.split(']:', 1)
                    if len(parts) == 2:
                        persona = parts[0] + ']:'
                        fala = parts[1].strip()
                        
                        p = doc.add_paragraph()
                        p.add_run(persona).bold = True
                        p.add_run(f' {fala}')
                    else:
                        doc.add_paragraph(line)
                else:
                    doc.add_paragraph(line)
        
        if i < len(transcriptions):
            doc.add_paragraph("─" * 50)
    
    doc.save(output_path)
    print(f"Documento salvo em: {output_path}")

def process_transcriptions():
    """Processo principal de transcrição"""
    mp3_dir = "mp3"
    transcricao_dir = "trascricao"
    
    # Criar diretório de transcrição se não existir
    os.makedirs(transcricao_dir, exist_ok=True)
    
    # Encontrar todos os arquivos MP3 gerados
    mp3_files = sorted(glob.glob(os.path.join(mp3_dir, "*.mp3")))
    
    if not mp3_files:
        print("Nenhum arquivo MP3 encontrado no diretório mp3/")
        return
    
    print(f"Encontrados {len(mp3_files)} arquivos de áudio para transcrever")
    
    try:
        # Etapa 1: Identificar personas no primeiro arquivo
        first_file = mp3_files[0]
        personas_info = identify_personas(first_file)
        
        print("Personas identificadas:")
        for persona in personas_info['personas']:
            print(f"  - {persona['id']}: {persona['nome']}")
        
        # Etapa 2: Transcrever todos os arquivos
        transcriptions = []
        
        for i, mp3_file in enumerate(mp3_files, 1):
            print(f"\nProcessando arquivo {i}/{len(mp3_files)}")
            transcription = transcribe_audio_with_personas(mp3_file, personas_info)
            filename = os.path.basename(mp3_file)
            transcriptions.append((filename, transcription))
            
            # Pequena pausa entre requests para não sobrecarregar a API
            time.sleep(3)
        
        # Etapa 3: Criar documento Word
        base_name = Path(mp3_files[0]).stem.replace('-01', '')
        output_path = os.path.join(transcricao_dir, f"{base_name}_transcricao_completa.docx")
        
        create_transcription_document(transcriptions, personas_info, output_path)
        
        print(f"\n✅ Transcrição completa finalizada!")
        print(f"📄 Documento salvo em: {output_path}")
        
    except Exception as e:
        print(f"❌ Erro durante a transcrição: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    try:
        process_transcriptions()
    except KeyboardInterrupt:
        print("\n🛑 Processo interrompido pelo usuário")
    except Exception as e:
        print(f"❌ Erro fatal: {e}")
        import traceback
        traceback.print_exc()